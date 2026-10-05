import io
from typing import List
from datetime import datetime
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from app.domain.snapshot import SnapshotCache
from app.domain.metrics import calculate_period_kpi, compute_receivables
from app.schemas.domain import Project, Expense

class ReportService:
    def __init__(self, cache: SnapshotCache):
        self.cache = cache

    async def get_report_data(self, period_type: str, period_value: str):
        # period_type: 'month' or 'year'
        # period_value: 'YYYY-MM' or 'YYYY'
        projects, payments, expenses = await self.cache.get_data()

        def matches(bulan: str) -> bool:
            if period_type == "year":
                return bulan.startswith(period_value)
            return bulan == period_value

        filtered_projects = [p for p in projects if matches(p.bulan_filter)]
        filtered_expenses = [e for e in expenses if matches(e.bulan_filter)]
        kpi = calculate_period_kpi(projects, payments, expenses, period_type, period_value)

        # Receivables (global, tidak difilter periode) dari sumber tunggal
        receivables = [
            {
                "id": proj.id_proyek,
                "klien": proj.nama_klien,
                "tanggal": proj.tanggal,
                "nilai": proj.nilai_proyek,
                "sisa": sisa,
            }
            for proj, _, sisa in compute_receivables(projects, payments)
        ]

        return kpi, filtered_projects, filtered_expenses, receivables

    def generate_excel(self, kpi: dict, projects: List[Project], expenses: List[Expense], receivables: List[dict]) -> bytes:
        wb = Workbook()
        
        # 1. Sheet Ringkasan
        ws1 = wb.active
        ws1.title = "Ringkasan KPI"
        ws1.append(["Metrik", "Nilai (Rp)"])
        ws1.append(["Omzet", kpi["omzet"]["value"]])
        ws1.append(["Kas Masuk", kpi["kas_masuk"]["value"]])
        ws1.append(["Pengeluaran", kpi["pengeluaran"]["value"]])
        ws1.append(["Laba Bersih", kpi["laba_bersih"]["value"]])
        
        # 2. Sheet Proyek
        ws2 = wb.create_sheet("Proyek")
        ws2.append(["No Invoice", "Tanggal", "Klien", "Pekerjaan", "Nilai Proyek", "Status"])
        for p in projects:
            ws2.append([p.id_proyek, p.tanggal.isoformat(), p.nama_klien, p.pekerjaan, p.nilai_proyek, p.status_bayar])
            
        # 3. Sheet Pengeluaran
        ws3 = wb.create_sheet("Pengeluaran")
        ws3.append(["ID", "Tanggal", "Kategori", "Keterangan", "Nominal"])
        for e in expenses:
            ws3.append([e.id_pengeluaran, e.tanggal.isoformat(), e.kategori, e.keterangan, e.nominal])
            
        # 4. Sheet Piutang
        ws4 = wb.create_sheet("Piutang")
        ws4.append(["No Invoice", "Tanggal", "Klien", "Nilai Proyek", "Sisa Tagihan"])
        for r in receivables:
            ws4.append([r["id"], r["tanggal"].isoformat(), r["klien"], r["nilai"], r["sisa"]])
            
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    def generate_pdf(self, title: str, kpi: dict, projects: List[Project], expenses: List[Expense], receivables: List[dict]) -> bytes:
        out = io.BytesIO()
        doc = SimpleDocTemplate(out, pagesize=landscape(A4))
        elements = []
        styles = getSampleStyleSheet()
        
        elements.append(Paragraph(title, styles['Title']))
        elements.append(Spacer(1, 20))
        
        # KPI Table
        elements.append(Paragraph("Ringkasan Keuangan", styles['Heading2']))
        kpi_data = [
            ["Metrik", "Nilai (Rp)"],
            ["Omzet", f"{kpi['omzet']['value']:,}"],
            ["Kas Masuk", f"{kpi['kas_masuk']['value']:,}"],
            ["Pengeluaran", f"{kpi['pengeluaran']['value']:,}"],
            ["Laba Bersih", f"{kpi['laba_bersih']['value']:,}"]
        ]
        t = Table(kpi_data, colWidths=[150, 150])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.grey),
            ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('GRID', (0,0), (-1,-1), 1, colors.black)
        ]))
        elements.append(t)
        elements.append(Spacer(1, 20))
        
        # Piutang Table
        elements.append(Paragraph("Daftar Piutang Berjalan", styles['Heading2']))
        if receivables:
            piutang_data = [["No Invoice", "Klien", "Tanggal", "Nilai Proyek", "Sisa"]]
            for r in receivables:
                piutang_data.append([r['id'], r['klien'], r['tanggal'].isoformat(), f"{r['nilai']:,}", f"{r['sisa']:,}"])
            t2 = Table(piutang_data)
            t2.setStyle(TableStyle([('GRID', (0,0), (-1,-1), 1, colors.black), ('BACKGROUND', (0,0), (-1,0), colors.lightgrey)]))
            elements.append(t2)
        else:
            elements.append(Paragraph("Tidak ada piutang saat ini.", styles['Normal']))
            
        doc.build(elements)
        return out.getvalue()
