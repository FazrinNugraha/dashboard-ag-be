"""Pembuat Excel (XLSX) laporan keuangan dengan tampilan clean & soft blue.

Desain mengadopsi format laporan posisi keuangan modern:
- Kop laporan ringkas 3 baris di atas (Judul, Nama Perusahaan, Periode)
- Header tabel biru corporate elegan (Calibri, teks putih bold)
- Banner pembatas kategori berwarna biru muda lembut (soft blue)
- Baris data putih bersih dengan gridlines aktif dan garis kotak-kotak rapi
- Baris total dengan garis bawah ganda (accounting underline) dan highlight biru lembut
- Status badge LUNAS (hijau soft) dan DP (kuning amber soft)
"""
import io
from datetime import date
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.domain import Expense, Project
from app.services.report_pdf import CATEGORY_LABELS, COMPANY, format_date

# --- Palet Warna (Format openpyxl ARGB hex tanpa '#') ------------------------
PRIMARY_BLUE = "1D4ED8"        # Biru Corporate Solid (Header Tabel Utama)
SECTION_BLUE = "DBEAFE"        # Biru Muda Lembut (Banner Kategori)
SECTION_FG = "1E40AF"          # Biru Tua Gelap untuk teks banner kategori
TOTAL_BG = "DBEAFE"            # Biru lembut untuk baris total / subtotal
STATUS_LUNAS_BG = "DCFCE7"
STATUS_LUNAS_FG = "166534"
STATUS_DP_BG = "FEF3C7"
STATUS_DP_FG = "92400E"

# Format angka Excel. Laba: hijau bila untung, merah bila rugi.
MONEY_FMT = '"Rp"#,##0'
PROFIT_FMT = '[Color10]"Rp"#,##0;[Red]-"Rp"#,##0'
PCT_FMT = "0.0%"
DATE_FMT = "dd/mm/yyyy"

FONT_FAMILY = "Calibri"
HEADER_ROW = 5

_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_TOTAL_BORDER = Border(
    left=_THIN,
    right=_THIN,
    top=_THIN,
    bottom=Side(style="double", color="000000"),
)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


def _letterhead(ws: Worksheet, ncols: int, title: str, label: str) -> None:
    """Kop laporan ringkas & bersih (3 baris): Judul Laporan, Nama Perusahaan, Periode."""
    def write_centered(row: int, text: str, font: Font, height: float = 20):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = height

    write_centered(1, title, Font(name=FONT_FAMILY, size=13, bold=True, color=PRIMARY_BLUE), height=24)
    write_centered(2, COMPANY["name"], Font(name=FONT_FAMILY, size=11, bold=True, color="1E293B"), height=18)
    if label:
        write_centered(3, f"Periode: {label}", Font(name=FONT_FAMILY, size=10, color="64748B"), height=16)
    ws.row_dimensions[4].height = 8  # Spacer row


def _table(
    ws: Worksheet,
    headers: List[str],
    rows: List[list],
    widths: List[float],
    money: tuple = (),
    profit: tuple = (),
    pct: tuple = (),
    dates: tuple = (),
    center: tuple = (),
    total: Optional[list] = None,
) -> int:
    """Tulis tabel mulai HEADER_ROW. Kolom 1-based. Mengembalikan baris terakhir."""
    ncols = len(headers)

    # Header Biru Corporate
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=HEADER_ROW, column=c, value=h)
        cell.font = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
        cell.fill = _fill(PRIMARY_BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
    ws.row_dimensions[HEADER_ROW].height = 24

    def style_cell(cell, c: int, bold: bool = False, is_total: bool = False):
        cell.font = Font(name=FONT_FAMILY, size=10, bold=bold)
        cell.border = _TOTAL_BORDER if is_total else _BORDER
        cell.alignment = Alignment(
            horizontal="center" if c in center else "right" if c in money + profit + pct else "left",
            vertical="center",
            wrap_text=True,
        )
        if c in money:
            cell.number_format = MONEY_FMT
        elif c in profit:
            cell.number_format = PROFIT_FMT
        elif c in pct:
            cell.number_format = PCT_FMT
        elif c in dates:
            cell.number_format = DATE_FMT

    # Baris data putih bersih
    for r, values in enumerate(rows):
        row_idx = HEADER_ROW + 1 + r
        ws.row_dimensions[row_idx].height = 20
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=c, value=v)
            style_cell(cell, c)
            cell.fill = _fill("FFFFFF")

    last = HEADER_ROW + len(rows)

    # Baris total
    if total:
        last += 1
        ws.row_dimensions[last].height = 22
        for c, v in enumerate(total, 1):
            cell = ws.cell(row=last, column=c, value=v)
            style_cell(cell, c, bold=True, is_total=True)
            cell.fill = _fill(TOTAL_BG)

    # Lebar kolom, bekukan header, filter, dan pengaturan cetak/tampilan
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)
    if rows:
        data_end = HEADER_ROW + len(rows)
        ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(ncols)}{data_end}"
    ws.sheet_view.showGridLines = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    return last


def _footer_note(ws: Worksheet, ncols: int, last_row: int, printed_on: date) -> None:
    """Catatan kecil di bawah tabel: tanggal cetak."""
    ws.merge_cells(start_row=last_row + 2, start_column=1, end_row=last_row + 2, end_column=ncols)
    cell = ws.cell(row=last_row + 2, column=1, value=f"Dicetak pada {format_date(printed_on)}")
    cell.font = Font(name=FONT_FAMILY, size=9, italic=True, color="64748B")
    cell.alignment = Alignment(horizontal="right")


def _status_fill(ws: Worksheet, col: int, first_row: int, last_row: int) -> None:
    """Warnai sel status: LUNAS hijau soft, DP kuning soft."""
    for r in range(first_row, last_row + 1):
        cell = ws.cell(row=r, column=col)
        if cell.value == "LUNAS":
            cell.fill = _fill(STATUS_LUNAS_BG)
            cell.font = Font(name=FONT_FAMILY, size=10, bold=True, color=STATUS_LUNAS_FG)
        elif cell.value == "DP":
            cell.fill = _fill(STATUS_DP_BG)
            cell.font = Font(name=FONT_FAMILY, size=10, bold=True, color=STATUS_DP_FG)


def build_report_excel(
    kpi: dict,
    projects: List[Project],
    expenses: List[Expense],
    receivables: List[dict],
    label: str,
    printed_on: date,
    project_expenses: Optional[Dict[str, int]] = None,
) -> bytes:
    """Bangun workbook laporan (4 sheet) dengan tema clean & soft blue."""
    if project_expenses is None:
        project_expenses = {}
        for e in expenses:
            if e.id_proyek:
                project_expenses[e.id_proyek] = project_expenses.get(e.id_proyek, 0) + e.nominal

    wb = Workbook()

    # 1. Ringkasan KPI & Posisi Keuangan -------------------------------------
    ws = wb.active
    ws.title = "Ringkasan KPI"
    ws.sheet_properties.tabColor = PRIMARY_BLUE
    margin = kpi["laba_bersih"].get("margin_pct")
    _letterhead(ws, 2, "LAPORAN RINGKASAN KEUANGAN", label)

    # Header
    for c, h in enumerate(["Komponen Keuangan", "Nominal"], 1):
        cell = ws.cell(row=HEADER_ROW, column=c, value=h)
        cell.font = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
        cell.fill = _fill(PRIMARY_BLUE)
        cell.alignment = Alignment(horizontal="center" if c == 2 else "left", vertical="center")
        cell.border = _BORDER
    ws.row_dimensions[HEADER_ROW].height = 24

    curr_row = HEADER_ROW + 1

    def add_section_banner(title: str):
        nonlocal curr_row
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=2)
        cell = ws.cell(row=curr_row, column=1, value=title)
        cell.font = Font(name=FONT_FAMILY, size=10, bold=True, color=SECTION_FG)
        cell.fill = _fill(SECTION_BLUE)
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.cell(row=curr_row, column=2).fill = _fill(SECTION_BLUE)
        for col in (1, 2):
            ws.cell(row=curr_row, column=col).border = _BORDER
        ws.row_dimensions[curr_row].height = 22
        curr_row += 1

    def add_data_row(label_text: str, value, num_fmt: str = MONEY_FMT, bold: bool = False, is_subtotal: bool = False):
        nonlocal curr_row
        c1 = ws.cell(row=curr_row, column=1, value=label_text)
        c2 = ws.cell(row=curr_row, column=2, value=value)
        c1.font = Font(name=FONT_FAMILY, size=10, bold=bold)
        c2.font = Font(name=FONT_FAMILY, size=10, bold=bold)
        c1.alignment = Alignment(horizontal="left", vertical="center", indent=2 if not is_subtotal else 1)
        c2.alignment = Alignment(horizontal="right", vertical="center")
        c2.number_format = num_fmt
        bg = TOTAL_BG if is_subtotal else "FFFFFF"
        c1.fill = _fill(bg)
        c2.fill = _fill(bg)
        border = _TOTAL_BORDER if is_subtotal else _BORDER
        c1.border = border
        c2.border = border
        ws.row_dimensions[curr_row].height = 20
        curr_row += 1

    # Section 1: ARUS KAS & PENDAPATAN
    add_section_banner("ARUS KAS & PENDAPATAN")
    add_data_row("Omzet (nilai invoice terbit)", kpi["omzet"]["value"])
    add_data_row("Kas Masuk (pembayaran diterima)", kpi["kas_masuk"]["value"])

    # Section 2: PENGELUARAN (BEBAN)
    add_section_banner("PENGELUARAN (BEBAN OPERASIONAL & PROYEK)")
    category_totals = {k: 0 for k in CATEGORY_LABELS}
    for e in expenses:
        cat = e.kategori if e.kategori in category_totals else "LAINNYA"
        category_totals[cat] += e.nominal

    for cat_key, cat_label in CATEGORY_LABELS.items():
        amt = category_totals.get(cat_key, 0)
        if amt > 0:
            add_data_row(f"{cat_label}", amt)

    add_data_row("Total Pengeluaran", kpi["pengeluaran"]["value"], is_subtotal=True, bold=True)

    # Section 3: HASIL BERSIH
    add_section_banner("HASIL BERSIH")
    add_data_row(
        "Laba Bersih (Kas Masuk - Pengeluaran)",
        kpi["laba_bersih"]["value"],
        num_fmt=PROFIT_FMT,
        bold=True,
        is_subtotal=True,
    )
    if margin is not None:
        add_data_row("Margin Laba Bersih", margin / 100, num_fmt=PCT_FMT)

    # Section 4: PIUTANG
    if receivables:
        add_section_banner("PIUTANG BERJALAN")
        total_piutang = sum(r["sisa"] for r in receivables)
        add_data_row("Total Sisa Tagihan (Piutang)", total_piutang, is_subtotal=True, bold=True)

    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 26
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)
    ws.sheet_view.showGridLines = True
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    _footer_note(ws, 2, curr_row - 1, printed_on)

    # 2. Proyek (dengan laba per proyek) ------------------------------------
    ws = wb.create_sheet("Proyek")
    ws.sheet_properties.tabColor = "1D4ED8"
    _letterhead(ws, 9, "LAPORAN PROYEK DAN LABA PER PROYEK", label)
    rows, t_nilai, t_keluar = [], 0, 0
    for p in sorted(projects, key=lambda x: x.tanggal):
        keluar = project_expenses.get(p.id_proyek, 0)
        laba = p.nilai_proyek - keluar
        rows.append([
            p.id_proyek, p.tanggal, p.nama_klien, p.pekerjaan, p.nilai_proyek, keluar, laba,
            laba / p.nilai_proyek if p.nilai_proyek > 0 else "-", p.status_bayar,
        ])
        t_nilai += p.nilai_proyek
        t_keluar += keluar
    last = _table(
        ws, ["No. Invoice", "Tanggal", "Klien", "Pekerjaan", "Nilai Proyek",
             "Pengeluaran", "Laba Bersih", "Margin", "Status"],
        rows, [16, 13, 22, 34, 18, 18, 18, 11, 12],
        money=(5, 6), profit=(7,), pct=(8,), dates=(2,), center=(1, 2, 9),
        total=["", "", "", "TOTAL", t_nilai, t_keluar, t_nilai - t_keluar, "", ""] if rows else None,
    )
    _status_fill(ws, 9, HEADER_ROW + 1, HEADER_ROW + len(rows))
    _footer_note(ws, 9, last, printed_on)

    # 3. Pengeluaran ---------------------------------------------------------
    ws = wb.create_sheet("Pengeluaran")
    ws.sheet_properties.tabColor = "0284C7"
    _letterhead(ws, 6, "LAPORAN PENGELUARAN", label)
    rows = [
        [e.id_pengeluaran, e.tanggal, CATEGORY_LABELS.get(e.kategori, e.kategori),
         e.keterangan, e.id_proyek or "Umum", e.nominal]
        for e in sorted(expenses, key=lambda x: x.tanggal)
    ]
    last = _table(
        ws, ["ID Pengeluaran", "Tanggal", "Kategori", "Keterangan", "Proyek", "Nominal"],
        rows, [16, 13, 18, 40, 18, 18], money=(6,), dates=(2,), center=(1, 2, 5),
        total=["", "", "", "", "TOTAL PENGELUARAN", sum(e.nominal for e in expenses)] if rows else None,
    )
    _footer_note(ws, 6, last, printed_on)

    # 4. Piutang -------------------------------------------------------------
    ws = wb.create_sheet("Piutang")
    ws.sheet_properties.tabColor = "0EA5E9"
    _letterhead(ws, 5, "DAFTAR PIUTANG BERJALAN", label)
    rows = [[r["id"], r["tanggal"], r["klien"], r["nilai"], r["sisa"]] for r in receivables]
    last = _table(
        ws, ["No. Invoice", "Tanggal", "Klien", "Nilai Proyek", "Sisa Tagihan"],
        rows, [16, 13, 26, 20, 20], money=(4, 5), dates=(2,), center=(1, 2),
        total=["", "", "TOTAL PIUTANG", sum(r["nilai"] for r in receivables), sum(r["sisa"] for r in receivables)] if rows else None,
    )
    _footer_note(ws, 5, last, printed_on)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
