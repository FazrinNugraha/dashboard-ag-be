"""Test ReportService: pengambilan data, Excel, dan PDF (tanpa Sheets)."""
import asyncio
from datetime import date

from app.domain.snapshot import SnapshotCache
from app.schemas.domain import Expense, Payment, Project
from app.services.report_service import ReportService


class InMemoryRepository:
    def __init__(self, projects, payments, expenses):
        self._data = (list(projects), list(payments), list(expenses))

    async def read_all(self):
        p, pay, e = self._data
        return list(p), list(pay), list(e)


def project(id_proyek, nilai, bulan):
    return Project(
        id_proyek=id_proyek,
        tanggal=date(int(bulan[:4]), int(bulan[5:7]), 1),
        nama_klien="Klien",
        alamat="A",
        pekerjaan="P",
        subtotal=nilai,
        diskon=0,
        nilai_proyek=nilai,
        bulan_filter=bulan,
        dibuat_oleh="admin1",
    )


def payment(id_bayar, id_proyek, nominal, bulan):
    return Payment(
        id_bayar=id_bayar,
        id_proyek=id_proyek,
        tanggal=date(int(bulan[:4]), int(bulan[5:7]), 1),
        nominal=nominal,
        tipe="DP",
        bulan_filter=bulan,
        dicatat_oleh="admin1",
    )


def expense(nominal, bulan):
    return Expense(
        id_pengeluaran="OUT-1",
        tanggal=date(int(bulan[:4]), int(bulan[5:7]), 1),
        kategori="UPAH",
        keterangan="x",
        nominal=nominal,
        bulan_filter=bulan,
        dibuat_oleh="admin1",
    )


def build_service():
    projects = [
        project("INV-2605-001", 10_000_000, "2026-05"),
        project("INV-2606-001", 20_000_000, "2026-06"),
    ]
    payments = [
        payment("PAY-1", "INV-2605-001", 4_000_000, "2026-05"),
        payment("PAY-2", "INV-2606-001", 5_000_000, "2026-06"),
    ]
    expenses = [expense(1_000_000, "2026-06")]
    repo = InMemoryRepository(projects, payments, expenses)
    cache = SnapshotCache(repo, ttl_seconds=30)
    asyncio.run(cache.force_refresh())
    return ReportService(cache)


class TestGetReportData:
    def test_periode_bulan(self):
        service = build_service()
        kpi, projects, expenses, receivables = asyncio.run(
            service.get_report_data("month", "2026-06")
        )
        assert kpi["omzet"]["value"] == 20_000_000
        assert kpi["kas_masuk"]["value"] == 5_000_000
        assert len(projects) == 1
        assert len(expenses) == 1
        # piutang global: INV-2605 sisa 6jt + INV-2606 sisa 15jt
        assert sum(r["sisa"] for r in receivables) == 21_000_000

    def test_periode_tahun(self):
        service = build_service()
        kpi, projects, expenses, _ = asyncio.run(service.get_report_data("year", "2026"))
        assert kpi["omzet"]["value"] == 30_000_000
        assert kpi["kas_masuk"]["value"] == 9_000_000
        assert len(projects) == 2

    def test_tidak_ada_data(self):
        service = build_service()
        kpi, projects, _, receivables = asyncio.run(service.get_report_data("month", "2030-01"))
        assert kpi["omzet"]["value"] == 0
        assert projects == []
        # piutang tetap global
        assert receivables


class TestGenerateExcel:
    def test_menghasilkan_xlsx_valid(self):
        service = build_service()
        kpi, projects, expenses, receivables = asyncio.run(
            service.get_report_data("month", "2026-06")
        )
        content = service.generate_excel(kpi, projects, expenses, receivables)
        assert content[:2] == b"PK"  # signature zip/xlsx
        assert len(content) > 1000

        from openpyxl import load_workbook
        import io

        wb = load_workbook(io.BytesIO(content))
        assert wb.sheetnames == ["Ringkasan KPI", "Proyek", "Pengeluaran", "Piutang"]


class TestGeneratePdf:
    def test_menghasilkan_pdf_valid(self):
        service = build_service()
        kpi, projects, expenses, receivables = asyncio.run(
            service.get_report_data("month", "2026-06")
        )
        content = service.generate_pdf("Laporan", kpi, projects, expenses, receivables)
        assert content.startswith(b"%PDF")